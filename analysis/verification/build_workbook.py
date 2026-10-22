"""Build the verification workbook for the small fake grid (README.md, step 3).

    python3 build_workbook.py [FOLDER]

FOLDER (default: this script's folder) holds what make_grid.py and exact_p.py wrote:
analysis-output/, planted-truth.csv, bundle-for-workbook.json and exact-p.csv. The workbook is
written to FOLDER/analysis-verification.xlsx. Needs openpyxl. Excel computes the formulas when
the workbook is opened.

Inputs (blue in the workbook) are copied from the analysis output, the fake grid's files and
the planted values. Every other number is an Excel formula.
"""
import csv
import json
import sys
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

import os
U = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(U, "analysis-verification.xlsx")

BLUE = Font(name="Arial", size=10, color="0000FF")
BLACK = Font(name="Arial", size=10, color="000000")
BOLD = Font(name="Arial", size=10, bold=True)
TITLE = Font(name="Arial", size=14, bold=True)
HEAD_FILL = PatternFill("solid", start_color="D9E1F2")
IN_FILL = PatternFill("solid", start_color="EEF3FB")
OK_FILL = PatternFill("solid", start_color="C6EFCE")
BAD_FILL = PatternFill("solid", start_color="FFC7CE")
WRAP = Alignment(wrap_text=True, vertical="top")


def rd(path):
    with open(path, encoding="utf-8") as f:
        return list(csv.DictReader(f))


runs = rd(f"{U}/analysis-output/run-table.csv")
conf = rd(f"{U}/analysis-output/confirmatory.csv")
sens = rd(f"{U}/analysis-output/sensitivity.csv")
voids = rd(f"{U}/analysis-output/voids.csv")
planted = {r["blind_id"]: r for r in rd(f"{U}/planted-truth.csv")}
exact = rd(f"{U}/exact-p.csv")
bundle = json.load(open(f"{U}/bundle-for-workbook.json"))

runs.sort(key=lambda r: (r["model"], r["task"], r["condition"], ["off", "unseen", "seen"].index(r["arm"]), int(r["replicate"])))
N = len(runs)
R1, RN = 2, N + 1


def num(x):
    if x in ("", "nan", None):
        return None
    try:
        return float(x)
    except ValueError:
        return x


def put(ws, cell, value, font=BLACK, fill=None, fmt=None):
    c = ws[cell]
    c.value = value
    c.font = font
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def header(ws, row, names, widths=None):
    for i, n in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=n)
        c.font = BOLD
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[L(i)].width = w


def status_format(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"ok"'], fill=OK_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"close"'], fill=OK_FILL))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"CHECK"'], fill=BAD_FILL))


wb = Workbook()
readme = wb.active
readme.title = "Read me"
settings = wb.create_sheet("Settings")
ws = wb.create_sheet("Runs")
qs = wb.create_sheet("Quality")
vs = wb.create_sheet("Voids")
ts = wb.create_sheet("Tests")

# ------------------------------------------------------------------ Settings
header(settings, 1, ["setting", "value", "where it comes from"], [34, 12, 70])
SET = [
    ("alpha", 0.05, "protocol section 8.1, Holm at alpha = 0.05"),
    ("failure equivalence bound (+/-)", 0.10, "protocol section 8.2, +/-10 percentage points"),
    ("steps equivalence bound, low", 0.8, "protocol section 8.2"),
    ("steps equivalence bound, high", 1.25, "protocol section 8.2"),
    ("permutations", 10000, "protocol section 8.1 (the Monte Carlo p is (1 + hits) / (10,000 + 1))"),
    ("z for a 95% interval", "=NORMSINV(0.975)", "formula"),
    ("z for a 90% interval", "=NORMSINV(0.95)", "formula"),
    ("tolerance for 6-digit CSV numbers", 1e-5, "the analysis CSVs print 6 significant digits, so matches are checked to a relative 1e-5"),
    ("interval 'close' tolerance", 0.25, "a bootstrap interval end counts as close to the formula interval's end if they differ by at most this share of the interval's width"),
]
for i, (k, v, src) in enumerate(SET, 2):
    put(settings, f"A{i}", k)
    put(settings, f"B{i}", v, BLACK if isinstance(v, str) else BLUE)
    put(settings, f"C{i}", src)
S_ALPHA, S_FB, S_SLO, S_SHI, S_PERM, S_Z95, S_Z90, S_TOL, S_CLOSE = [f"Settings!$B${i}" for i in range(2, 11)]

# ------------------------------------------------------------------ Runs
RCOL = {}
names = [
    ("blind_id", "in"), ("model", "in"), ("task", "in"), ("condition", "in"), ("arm", "in"), ("replicate", "in"),
    ("status", "in"), ("steps", "in"), ("nudged", "in"), ("reasoning_rate", "in"), ("reasoning_median_chars", "in"),
    ("silent_turn", "in"), ("sampled", "in"),
    ("stratum", "f"), ("failed", "f"), ("log_steps", "f"), ("log_median_chars", "f"), ("quality", "f"),
    ("quality, substituted scored 0", "f"), ("substituted", "f"), ("api_conformance", "f"),
    ("analysis: failed", "in"), ("analysis: log_steps", "in"), ("analysis: log median", "in"),
    ("analysis: quality", "in"), ("analysis: substituted", "in"), ("analysis: api_conformance", "in"),
    ("planted: status", "in"), ("planted: steps", "in"), ("planted: reasoning_rate", "in"),
    ("planted: median chars", "in"), ("planted: silent_turn", "in"), ("planted: nudged", "in"),
    ("check: analysis read the files right (vs planted)", "f"),
    ("check: Excel recompute = analysis", "f"),
]
for i, (n, _) in enumerate(names, 1):
    RCOL[n] = L(i)
helper_measures = [("steps", "log_steps"), ("quality", "quality"), ("H3", "reasoning_rate"), ("H4", "log_median_chars")]
col = len(names) + 2
for m, src in helper_measures:
    for part in ("has", "k in stratum", "stratum mean", "w"):
        RCOL[f"{m}:{part}"] = L(col)
        col += 1
allnames = [n for n, _ in names] + [""] + [f"{m}: {p}" for m, _ in helper_measures for p in ("has value", "k in stratum", "stratum mean", "w (for the variance)")]
header(ws, 1, allnames)
for i in range(1, len(allnames) + 1):
    ws.column_dimensions[L(i)].width = 11
ws.column_dimensions["A"].width = 9
ws.column_dimensions["B"].width = 11
ws.column_dimensions["C"].width = 13
ws.column_dimensions["D"].width = 15
ws.row_dimensions[1].height = 75
ws.freeze_panes = "B2"


def rng(name):
    c = RCOL[name]
    return f"Runs!${c}${R1}:${c}${RN}"


C = RCOL
for j, r in enumerate(runs):
    i = R1 + j
    p = planted[r["blind_id"]]
    vals = {"blind_id": r["blind_id"], "model": r["model"], "task": r["task"], "condition": r["condition"],
            "arm": r["arm"], "replicate": int(r["replicate"]), "status": r["status"], "steps": num(r["steps"]),
            "nudged": num(r["nudged"]), "reasoning_rate": num(r["reasoning_rate"]),
            "reasoning_median_chars": num(r["reasoning_median_chars"]), "silent_turn": num(r["silent_turn"]),
            "sampled": num(r["sampled"]),
            "analysis: failed": num(r["failed"]), "analysis: log_steps": num(r["log_steps"]),
            "analysis: log median": num(r["log_reasoning_median_chars"]), "analysis: quality": num(r["quality"]),
            "analysis: substituted": num(r["substituted"]), "analysis: api_conformance": num(r["api_conformance"]),
            "planted: status": p["status"], "planted: steps": num(p["steps"]), "planted: reasoning_rate": num(p["reasoning_rate"]),
            "planted: median chars": num(p["reasoning_median_chars"]), "planted: silent_turn": num(p["silent_turn"]),
            "planted: nudged": num(p["nudged"])}
    for k, v in vals.items():
        put(ws, f"{C[k]}{i}", v, BLUE, IN_FILL)
    f = {
        "stratum": f'={C["task"]}{i}&"|"&{C["condition"]}{i}',
        "failed": f'=IF({C["status"]}{i}="finished",0,1)',
        "log_steps": f'=IF({C["failed"]}{i}=0,LN({C["steps"]}{i}),"")',
        "log_median_chars": f'=IF(ISNUMBER({C["reasoning_median_chars"]}{i}),IF({C["reasoning_median_chars"]}{i}>0,LN({C["reasoning_median_chars"]}{i}),""),"")',
        "quality": f'=IF({C["sampled"]}{i}=1,INDEX(Quality!$AF:$AF,MATCH({C["blind_id"]}{i},Quality!$A:$A,0)),"")',
        "substituted": f'=IF({C["sampled"]}{i}=1,INDEX(Quality!$AG:$AG,MATCH({C["blind_id"]}{i},Quality!$A:$A,0)),0)',
        "quality, substituted scored 0": f'=IF({C["quality"]}{i}="","",IF({C["substituted"]}{i}=1,0,{C["quality"]}{i}))',
        "api_conformance": f'=IF({C["sampled"]}{i}=1,INDEX(Quality!$AS:$AS,MATCH({C["blind_id"]}{i},Quality!$A:$A,0)),"")',
    }
    for k, v in f.items():
        put(ws, f"{C[k]}{i}", v)

    def eqn(a, b):
        return (f'IF(AND(ISNUMBER({a}),ISNUMBER({b})),ABS(N({a})-N({b}))<={S_TOL}*MAX(1,ABS(N({b}))),'
                f'AND({a}="",{b}=""))')

    put(ws, f'{C["check: analysis read the files right (vs planted)"]}{i}',
        f'=IF(AND({C["status"]}{i}={C["planted: status"]}{i},{eqn(C["steps"]+str(i), C["planted: steps"]+str(i))},'
        f'{eqn(C["reasoning_rate"]+str(i), C["planted: reasoning_rate"]+str(i))},'
        f'{eqn(C["reasoning_median_chars"]+str(i), C["planted: median chars"]+str(i))},'
        f'{C["silent_turn"]}{i}={C["planted: silent_turn"]}{i},{C["nudged"]}{i}={C["planted: nudged"]}{i}),"ok","CHECK")')
    put(ws, f'{C["check: Excel recompute = analysis"]}{i}',
        f'=IF(AND({eqn(C["failed"]+str(i), C["analysis: failed"]+str(i))},'
        f'{eqn(C["log_steps"]+str(i), C["analysis: log_steps"]+str(i))},'
        f'{eqn(C["log_median_chars"]+str(i), C["analysis: log median"]+str(i))},'
        f'{eqn(C["quality"]+str(i), C["analysis: quality"]+str(i))},'
        f'{C["substituted"]}{i}={C["analysis: substituted"]}{i},'
        f'{eqn(C["api_conformance"]+str(i), C["analysis: api_conformance"]+str(i))}),"ok","CHECK")')
    for m, src in helper_measures:
        v = f"{C[src]}{i}"
        has, k, mean, w = (C[f"{m}:{x}"] for x in ("has", "k in stratum", "stratum mean", "w"))
        crit = (f'{rng("model")},{C["model"]}{i},{rng("arm")},{C["arm"]}{i},{rng("stratum")},{C["stratum"]}{i},'
                f'Runs!${has}${R1}:${has}${RN},1')
        put(ws, f"{has}{i}", f"=IF(ISNUMBER({v}),1,0)")
        put(ws, f"{k}{i}", f"=COUNTIFS({crit})")
        put(ws, f"{mean}{i}", f'=IF({has}{i}=1,AVERAGEIFS(Runs!${C[src]}${R1}:${C[src]}${RN},{crit}),"")')
        put(ws, f"{w}{i}", f"=IF({has}{i}=1,IF({k}{i}>1,({v}-{mean}{i})^2*{k}{i}/({k}{i}-1),0),0)")
status_format(ws, f'{C["check: analysis read the files right (vs planted)"]}{R1}:{C["check: Excel recompute = analysis"]}{RN}')

# ------------------------------------------------------------------ Quality
cards = bundle["cards"]
qs["A1"] = "Quality score of every graded app, recomputed from its item results (score sheets). Score = 100 x (weight of items passed) / (weight of items passed or failed), rounded to 2 decimals, as score_sheet.py does."
qs["A1"].font = BOLD
qs.column_dimensions["A"].width = 9
qs.column_dimensions["B"].width = 13
qs.column_dimensions["C"].width = 10
qs.column_dimensions["D"].width = 22
for c in range(5, 47):
    qs.column_dimensions[L(c)].width = 7
for c in ("AF", "AG", "AS", "AT"):
    qs.column_dimensions[c].width = 12
row = 3
sample = {s["blind_id"] for s in bundle["sample"]}
for task in ("kanban-board", "message-board"):
    items = cards[task]
    put(qs, f"A{row}", f"{task}: item", BOLD, HEAD_FILL)
    put(qs, f"A{row + 1}", "weight", BOLD, HEAD_FILL)
    for j, it in enumerate(items):
        put(qs, f"{L(5 + j)}{row}", it["id"], BOLD, HEAD_FILL)
        put(qs, f"{L(5 + j)}{row + 1}", it["weight"], BLUE, IN_FILL)
    hdr = row + 2
    for c, n in (("A", "blind_id"), ("B", "task"), ("C", "accommodations"), ("D", "L0 notes"), ("AF", "score (Excel)"),
                 ("AG", "substituted (Excel)"), ("AS", "API conformance (Excel)"), ("AT", "in the analysis?")):
        put(qs, f"{c}{hdr}", n, BOLD, HEAD_FILL)
    for j in range(10):
        put(qs, f"{L(35 + j)}{hdr}", f"probe {j + 1}", BOLD, HEAD_FILL)
    wrow = row + 1
    r = hdr + 1
    for bid, sh in sorted(bundle["sheets"].items()):
        if sh["task"] != task:
            continue
        put(qs, f"A{r}", bid, BLUE, IN_FILL)
        put(qs, f"B{r}", task, BLUE, IN_FILL)
        put(qs, f"C{r}", " ".join(sh["acc"]), BLUE, IN_FILL)
        put(qs, f"D{r}", " ".join(sh["l0notes"].values()), BLUE, IN_FILL)
        for j, it in enumerate(items):
            put(qs, f"{L(5 + j)}{r}", sh["items"][it["id"]], BLUE, IN_FILL)
        for j, res in enumerate(bundle["probes"].get(bid, [])):
            put(qs, f"{L(35 + j)}{r}", res, BLUE, IN_FILL)
        rr = f"$E{r}:$AD{r}"
        ww = f"$E${wrow}:$AD${wrow}"
        put(qs, f"AF{r}", f'=ROUND(100*SUMPRODUCT(({rr}="pass")*{ww})/SUMPRODUCT((({rr}="pass")+({rr}="fail"))*{ww}),2)')
        put(qs, f"AG{r}", f'=IF(OR(ISNUMBER(SEARCH("A8",C{r})),AND(ISNUMBER(SEARCH("A0",C{r})),ISNUMBER(SEARCH("ETARGET",D{r})))),1,0)')
        pr = f"AI{r}:AR{r}"
        put(qs, f"AS{r}", f'=IF(COUNTIF({pr},"pass")+COUNTIF({pr},"fail")=0,"",COUNTIF({pr},"pass")/(COUNTIF({pr},"pass")+COUNTIF({pr},"fail")))')
        put(qs, f"AT{r}", f'=IF(COUNTIF({rng("blind_id")},A{r})>0,"yes","no, a void attempt")')
        r += 1
    row = r + 2
qs.freeze_panes = "E3"

# ------------------------------------------------------------------ Voids
header(vs, 1, ["blind_id", "label (the slot)", "model", "arm", "agent exit code", "status", "void? (Excel)",
               "in quality sample", "used in the analysis? (Excel)", "check: one analysed attempt per slot, never a void one"],
       [9, 44, 11, 8, 9, 20, 9, 9, 11, 18])
vs.row_dimensions[1].height = 60
key = sorted(bundle["key"], key=lambda k: (k["label"], k["started_at"]))
K1, KN = 2, len(key) + 1
for j, k in enumerate(key):
    i = K1 + j
    for c, v in (("A", k["blind_id"]), ("B", k["label"]), ("C", k["model_key"]), ("D", k["arm"]),
                 ("E", k["agent_rc"]), ("F", k["status"]), ("H", 1 if k["blind_id"] in sample else 0)):
        put(vs, f"{c}{i}", v, BLUE, IN_FILL)
    put(vs, f"G{i}", f'=IF(OR(E{i}<>0,F{i}="api_error"),1,0)')
    put(vs, f"I{i}", f'=IF(COUNTIF({rng("blind_id")},A{i})>0,1,0)')
    put(vs, f"J{i}", f'=IF(AND(COUNTIFS($B${K1}:$B${KN},B{i},$I${K1}:$I${KN},1)=1,NOT(AND(G{i}=1,I{i}=1))),"ok","CHECK")')
status_format(vs, f"J{K1}:J{KN}")
put(vs, "L1", "Per model x arm", BOLD)
hdrs = ["model", "arm", "attempts (Excel)", "void attempts (Excel)", "analysed runs (Excel)", "void attempts in the quality sample (Excel)",
        "analysis: attempts", "analysis: void attempts", "analysis: analysed runs", "analysis: void in quality sample", "check"]
for j, n in enumerate(hdrs):
    c = vs.cell(row=2, column=12 + j, value=n)
    c.font = BOLD
    c.fill = HEAD_FILL
    c.alignment = WRAP
    vs.column_dimensions[L(12 + j)].width = 12
vs.row_dimensions[2].height = 60
for j, v in enumerate(voids):
    i = 3 + j
    put(vs, f"L{i}", v["model"], BLUE, IN_FILL)
    put(vs, f"M{i}", v["arm"], BLUE, IN_FILL)
    put(vs, f"N{i}", f"=COUNTIFS($C${K1}:$C${KN},L{i},$D${K1}:$D${KN},M{i})")
    put(vs, f"O{i}", f"=COUNTIFS($C${K1}:$C${KN},L{i},$D${K1}:$D${KN},M{i},$G${K1}:$G${KN},1)")
    put(vs, f"P{i}", f"=COUNTIFS({rng('model')},L{i},{rng('arm')},M{i})")
    put(vs, f"Q{i}", f"=COUNTIFS($C${K1}:$C${KN},L{i},$D${K1}:$D${KN},M{i},$G${K1}:$G${KN},1,$H${K1}:$H${KN},1)")
    for c, k2 in (("R", "attempts"), ("S", "void_attempts"), ("T", "analysed_runs"), ("U", "void_in_quality_sample")):
        put(vs, f"{c}{i}", int(v[k2]), BLUE, IN_FILL)
    put(vs, f"V{i}", f'=IF(AND(N{i}=R{i},O{i}=S{i},P{i}=T{i},Q{i}=U{i}),"ok","CHECK")')
status_format(vs, f"V3:V{2 + len(voids)}")
VS_LAST = 2 + len(voids)

# ------------------------------------------------------------------ Tests
tests = conf + sens
ex = {(e["analysis"], e["model"], e["family"], e["measure"], e["arm_a"], e["arm_b"]): e for e in exact}
TH = ["analysis", "model", "family", "measure", "arm a", "arm b", "alternative", "what is compared",
      "n a (Excel)", "n b (Excel)", "mean a (Excel)", "mean b (Excel)", "estimate (Excel)",
      "analysis: n a", "analysis: n b", "analysis: mean a", "analysis: mean b", "analysis: estimate", "check: counts, means, estimate",
      "interval method", "95% low (Excel)", "95% high (Excel)", "analysis: 95% low", "analysis: 95% high",
      "90% low (Excel)", "90% high (Excel)", "analysis: 90% low", "analysis: 90% high", "check: intervals",
      "analysis: p (10,000 random relabellings)", "exact p (every relabelling, exact_p.py)", "relabellings counted", "check: p",
      "Holm p (Excel)", "analysis: Holm p", "check: Holm",
      "verdict (Excel)", "analysis: verdict", "check: verdict",
      "X (Excel)", "analysis: X", "check: X",
      "", "Wilson 95 a low", "Wilson 95 a high", "Wilson 95 b low", "Wilson 95 b high",
      "Wilson 90 a low", "Wilson 90 a high", "Wilson 90 b low", "Wilson 90 b high", "SE of the difference (bootstrap rows)"]
header(ts, 1, TH)
for i in range(1, len(TH) + 1):
    ts.column_dimensions[L(i)].width = 11
ts.column_dimensions["A"].width = 13
ts.column_dimensions["H"].width = 30
ts.row_dimensions[1].height = 75
ts.freeze_panes = "E2"
T = {n: L(i) for i, n in enumerate(TH, 1) if n}

WHAT = {"failure": ("failed", None, "share of runs not finished, all runs"),
        "steps": ("log_steps", "steps", "mean log steps, finished runs. Estimate = ratio of geometric means"),
        "quality": ("quality", "quality", "mean quality score, sampled apps"),
        "H3": ("reasoning_rate", "H3", "mean reasoning rate"),
        "H4": ("log_median_chars", "H4", "mean log of the run's median reasoning characters, runs with a reasoning step. Estimate = ratio"),
        "H5": ("silent_turn", None, "share of runs with a silent turn")}
row_of = {}
for j, t in enumerate(tests):
    row_of[(t["analysis"], t["model"], t["family"], t["measure"])] = 2 + j
for j, t in enumerate(tests):
    i = 2 + j
    m = t["measure"]
    src, helper, what = WHAT[m]
    if t["analysis"] == "substituted-0":
        src, helper, what = "quality, substituted scored 0", None, "mean quality score, sampled apps, substituted apps scored 0"
    if t["analysis"] == "no-nudge":
        what += ", runs without a nudge only"
    for c, v in (("analysis", t["analysis"]), ("model", t["model"]), ("family", t["family"]), ("measure", m),
                 ("arm a", t["arm_a"]), ("arm b", t["arm_b"]), ("alternative", t["alternative"] or "two-sided"),
                 ("what is compared", what)):
        put(ts, f"{T[c]}{i}", v, BLUE, IN_FILL)
    vr = rng(src)
    extra = f",{rng('nudged')},0" if t["analysis"] == "no-nudge" else ""
    hascrit = ""
    if src in ("log_steps", "quality", "log_median_chars", "quality, substituted scored 0"):
        hc = {"log_steps": "steps", "quality": "quality", "log_median_chars": "H4",
              "quality, substituted scored 0": "quality"}[src]
        hascrit = f",Runs!${C[hc + ':has']}${R1}:${C[hc + ':has']}${RN},1"

    def crit(arm_cell):
        return f"{rng('model')},$B{i},{rng('arm')},{arm_cell}{hascrit}{extra}"

    na, nb, ma, mb, est = (T[x] for x in ("n a (Excel)", "n b (Excel)", "mean a (Excel)", "mean b (Excel)", "estimate (Excel)"))
    put(ts, f"{na}{i}", f"=COUNTIFS({crit('$E' + str(i))})")
    put(ts, f"{nb}{i}", f"=COUNTIFS({crit('$F' + str(i))})")
    put(ts, f"{ma}{i}", f"=AVERAGEIFS({vr},{crit('$E' + str(i))})")
    put(ts, f"{mb}{i}", f"=AVERAGEIFS({vr},{crit('$F' + str(i))})")
    ratio = t["scale"] == "ratio"
    put(ts, f"{est}{i}", f"=EXP({mb}{i}-{ma}{i})" if ratio else f"={mb}{i}-{ma}{i}")
    for c, k2 in (("analysis: n a", "n_a"), ("analysis: n b", "n_b"), ("analysis: mean a", "mean_a"),
                  ("analysis: mean b", "mean_b"), ("analysis: estimate", "estimate"),
                  ("analysis: 95% low", "ci95_lo"), ("analysis: 95% high", "ci95_hi"), ("analysis: 90% low", "ci90_lo"),
                  ("analysis: 90% high", "ci90_hi"), ("analysis: p (10,000 random relabellings)", "p"),
                  ("analysis: Holm p", "p_holm"), ("analysis: X", "x")):
        put(ts, f"{T[c]}{i}", num(t[k2]), BLUE, IN_FILL)
    put(ts, f"{T['analysis: verdict']}{i}", t["verdict"], BLUE, IN_FILL)
    put(ts, f"{T['interval method']}{i}", t["interval"], BLUE, IN_FILL)

    def close(a, b):
        return f"ABS({a}-{b})<={S_TOL}*MAX(1,ABS({b}))"

    put(ts, f"{T['check: counts, means, estimate']}{i}",
        f'=IF(AND({na}{i}={T["analysis: n a"]}{i},{nb}{i}={T["analysis: n b"]}{i},'
        f'{close(ma + str(i), T["analysis: mean a"] + str(i))},{close(mb + str(i), T["analysis: mean b"] + str(i))},'
        f'{close(est + str(i), T["analysis: estimate"] + str(i))}),"ok","CHECK")')
    # Wilson pieces and Newcombe
    wa = [T[x] for x in ("Wilson 95 a low", "Wilson 95 a high", "Wilson 95 b low", "Wilson 95 b high",
                         "Wilson 90 a low", "Wilson 90 a high", "Wilson 90 b low", "Wilson 90 b high")]
    lo95, hi95, lo90, hi90 = (T[x] for x in ("95% low (Excel)", "95% high (Excel)", "90% low (Excel)", "90% high (Excel)"))
    se = T["SE of the difference (bootstrap rows)"]
    if t["interval"] == "newcombe":
        for z, cols in ((S_Z95, wa[:4]), (S_Z90, wa[4:])):
            for (p_, n_), (clo, chi) in (((ma, na), (cols[0], cols[1])), ((mb, nb), (cols[2], cols[3]))):
                p_c, n_c = f"{p_}{i}", f"{n_}{i}"
                core = f"({p_c}+{z}^2/(2*{n_c}))/(1+{z}^2/{n_c})"
                half = f"{z}*SQRT({p_c}*(1-{p_c})/{n_c}+{z}^2/(4*{n_c}^2))/(1+{z}^2/{n_c})"
                put(ts, f"{clo}{i}", f"=MAX(0,{core}-{half})")
                put(ts, f"{chi}{i}", f"=MIN(1,{core}+{half})")
        for (lo, hi), cols in (((lo95, hi95), wa[:4]), ((lo90, hi90), wa[4:])):
            la, ua, lb, ub = (f"{c}{i}" for c in cols)
            put(ts, f"{lo}{i}", f"=({mb}{i}-{ma}{i})-SQRT(({mb}{i}-{lb})^2+({ua}-{ma}{i})^2)")
            put(ts, f"{hi}{i}", f"=({mb}{i}-{ma}{i})+SQRT(({ub}-{mb}{i})^2+({ma}{i}-{la})^2)")
        exact_check = " AND ".join([])
        put(ts, f"{T['check: intervals']}{i}",
            f'=IF(AND({close(lo95 + str(i), T["analysis: 95% low"] + str(i))},{close(hi95 + str(i), T["analysis: 95% high"] + str(i))},'
            f'{close(lo90 + str(i), T["analysis: 90% low"] + str(i))},{close(hi90 + str(i), T["analysis: 90% high"] + str(i))}),"ok","CHECK")')
    elif helper is not None and t["analysis"] == "confirmatory":
        wcol = C[f"{helper}:w"]
        wr = f"Runs!${wcol}${R1}:${wcol}${RN}"
        put(ts, f"{se}{i}", f"=SQRT(SUMIFS({wr},{rng('model')},$B{i},{rng('arm')},$E{i})/{na}{i}^2+SUMIFS({wr},{rng('model')},$B{i},{rng('arm')},$F{i})/{nb}{i}^2)")
        d = f"({mb}{i}-{ma}{i})"
        for (lo, hi), z in (((lo95, hi95), S_Z95), ((lo90, hi90), S_Z90)):
            if ratio:
                put(ts, f"{lo}{i}", f"=EXP({d}-{z}*{se}{i})")
                put(ts, f"{hi}{i}", f"=EXP({d}+{z}*{se}{i})")
            else:
                put(ts, f"{lo}{i}", f"={d}-{z}*{se}{i}")
                put(ts, f"{hi}{i}", f"={d}+{z}*{se}{i}")
        a = {k: f"{T[k]}{i}" for k in ("analysis: 95% low", "analysis: 95% high", "analysis: 90% low", "analysis: 90% high")}
        w95 = f"({a['analysis: 95% high']}-{a['analysis: 95% low']})"
        w90 = f"({a['analysis: 90% high']}-{a['analysis: 90% low']})"
        put(ts, f"{T['check: intervals']}{i}",
            f'=IF(AND(ABS({lo95}{i}-{a["analysis: 95% low"]})<={S_CLOSE}*{w95},ABS({hi95}{i}-{a["analysis: 95% high"]})<={S_CLOSE}*{w95},'
            f'ABS({lo90}{i}-{a["analysis: 90% low"]})<={S_CLOSE}*{w90},ABS({hi90}{i}-{a["analysis: 90% high"]})<={S_CLOSE}*{w90}),"close","not close")')
    else:
        put(ts, f"{T['check: intervals']}{i}", "not checked (a subset)", BLACK)
    e = ex[(t["analysis"], t["model"], t["family"], m, t["arm_a"], t["arm_b"])]
    put(ts, f"{T['exact p (every relabelling, exact_p.py)']}{i}", float(e["exact_p"]), BLUE, IN_FILL)
    put(ts, f"{T['relabellings counted']}{i}", int(e["relabellings"]), BLUE, IN_FILL)
    pa, pe = T["analysis: p (10,000 random relabellings)"], T["exact p (every relabelling, exact_p.py)"]
    put(ts, f"{T['check: p']}{i}",
        f'=IF(ABS({pa}{i}-{pe}{i})<=3*SQRT({pe}{i}*(1-{pe}{i})/{S_PERM})+1/({S_PERM}+1),"ok","CHECK")')
    # Holm over the family: the three p values of this analysis, model and family. A
    # substituted-0 quality test is corrected with the headline family's failure and steps p.
    if t["analysis"] == "substituted-0":
        cells = [f"${pa}${row_of[('confirmatory', t['model'], t['family'], 'failure')]}",
                 f"${pa}${row_of[('confirmatory', t['model'], t['family'], 'steps')]}", f"{pa}{i}"]
    else:
        fam = [k for k in row_of if k[:3] == (t["analysis"], t["model"], t["family"])]
        cells = [f"${pa}${row_of[k]}" for k in fam]
    lo_ = f"MIN({','.join(cells)})"
    hi_ = f"MAX({','.join(cells)})"
    mid = f"({'+'.join(cells)}-{lo_}-{hi_})"
    own = f"{pa}{i}"
    put(ts, f"{T['Holm p (Excel)']}{i}",
        f"=MIN(1,IF({own}={lo_},3*{lo_},IF({own}<{hi_},MAX(3*{lo_},2*{own}),MAX(3*{lo_},2*{mid},{hi_}))))")
    put(ts, f"{T['check: Holm']}{i}", f'=IF({close(T["Holm p (Excel)"] + str(i), T["analysis: Holm p"] + str(i))},"ok","CHECK")')
    hp = f"{T['Holm p (Excel)']}{i}"
    if t["family"] == "H3-H5":
        v = f'=IF({hp}<{S_ALPHA},"supported","not supported")'
    elif m == "failure":
        v = f'=IF({hp}<{S_ALPHA},"different",IF(AND({lo90}{i}>-{S_FB},{hi90}{i}<{S_FB}),"equivalent","undetermined"))'
    elif m == "steps":
        v = (f'=IF({hp}<{S_ALPHA},"different",IF(AND({T["analysis: 90% low"]}{i}>{S_SLO},{T["analysis: 90% high"]}{i}<{S_SHI}),'
             f'"equivalent","undetermined"))')
    else:
        v = f'=IF({hp}<{S_ALPHA},"different","undetermined")'
    put(ts, f"{T['verdict (Excel)']}{i}", v)
    put(ts, f"{T['check: verdict']}{i}", f'=IF({T["verdict (Excel)"]}{i}={T["analysis: verdict"]}{i},"ok","CHECK")')
    if t["family"] == "H3-H5":
        put(ts, f"{T['X (Excel)']}{i}", "")
        put(ts, f"{T['check: X']}{i}", f'=IF({T["analysis: X"]}{i}="","ok","CHECK")')
    else:
        l9, h9 = f"{T['analysis: 90% low']}{i}", f"{T['analysis: 90% high']}{i}"
        xv = f"MAX({h9},1/{l9})" if ratio else f"MAX(ABS({l9}),ABS({h9}))"
        put(ts, f"{T['X (Excel)']}{i}", f'=IF({T["verdict (Excel)"]}{i}="different","",{xv})')
        put(ts, f"{T['check: X']}{i}",
            f'=IF(IF(AND(ISNUMBER({T["X (Excel)"]}{i}),ISNUMBER({T["analysis: X"]}{i})),'
            f'{close("N(" + T["X (Excel)"] + str(i) + ")", "N(" + T["analysis: X"] + str(i) + ")")},'
            f'AND({T["X (Excel)"]}{i}="",{T["analysis: X"]}{i}="")),"ok","CHECK")')
TN = 1 + len(tests)
for c in ("check: counts, means, estimate", "check: intervals", "check: p", "check: Holm", "check: verdict", "check: X"):
    status_format(ts, f"{T[c]}2:{T[c]}{TN}")
    ts.conditional_formatting.add(f"{T[c]}2:{T[c]}{TN}", CellIsRule(operator="equal", formula=['"not close"'], fill=BAD_FILL))
for r_ in range(2, TN + 1):
    for c in range(9, 53):
        cell = ts.cell(row=r_, column=c)
        if isinstance(cell.value, float) or (isinstance(cell.value, str) and cell.value.startswith("=") and "IF(" not in cell.value[:4]):
            cell.number_format = "0.0000"

# ------------------------------------------------------------------ Read me
readme.column_dimensions["A"].width = 30
readme.column_dimensions["B"].width = 110
put(readme, "A1", "Analysis verification on a small fake grid", TITLE)
lines = [
    ("What this is", "A check of the analysis code (plan/analysis) on a fake grid small enough to follow by hand. Every number the analysis produced is set beside the same number computed here by Excel formulas, from the run table's basic columns and the score sheets, or, for the permutation p values, by counting every possible relabelling."),
    ("The fake grid", f"Two models: gemma-q4 (three arms) and gptoss-20b (two arms, no OFF). Two runs in every task x condition cell (replicates r01 and r02 of the registered schedules), so 12 runs per arm. {len(bundle['key'])} attempts, {len(bundle['key']) - N} of them void, so {N} analysed runs. Every sampled app was graded with a filled score sheet and an API probe file. Made by make_grid.py (the seed and the planted values are written there)."),
    ("What was planted", "SEEN: 5% failure, about 40 steps, apps that work 85% of the time, reasoning on 98% of steps at about 150 characters, no silent turns. UNSEEN: 45% failure, about 70 steps, apps that work 35% of the time, reasoning on 70% of steps at about 400 characters, a silent turn in 40% of runs. OFF: 30% failure, no reasoning. Also 8% void attempts, 20% of runs nudged, 15% of graded apps with a substituted or added package. With 12 runs per arm, only the large effects reach significance, which is what a real test should do."),
    ("Colours", "Blue text on a light blue fill: a value copied in (from the analysis output, the fake grid's files, or the planted values). Black: an Excel formula. Green 'ok' or 'close': the check passes. Red 'CHECK' or 'not close': look at that row."),
    ("Runs", "One row per analysed run. Columns A to M are the run table's basic columns. N to U are recomputed by Excel (stratum, failed, log steps, log median, the quality score and API share looked up from the Quality sheet). Then the analysis's own values, then the planted values. Two checks per run: did the analysis read the files correctly (its values equal what was planted), and does Excel's recomputation equal the analysis. On the right, helper columns for the interval check."),
    ("Quality", "Every graded app's item results, the card's weights and the probe's check results. The score, the substituted flag and the API share are Excel formulas. Void attempts that fell in the quality sample are listed too, marked 'no, a void attempt'."),
    ("Voids", "Every attempt in the un-blinding key. Excel marks the void ones, checks that each slot has exactly one analysed attempt and that it is never a void one, and counts attempts and voids per model x arm against the analysis's voids table."),
    ("Tests", "One row per test of the analysis (confirmatory, and the no-nudge and substituted-0 sensitivity reruns). Counts, means and estimates are AVERAGEIFS and COUNTIFS on the Runs sheet. Newcombe intervals (failure, H5) are written out from Wilson's formula and must match exactly. Holm and the verdicts are formulas and must match exactly."),
    ("What Excel cannot redo", "Two numbers come from random resampling and cannot be repeated exactly by formulas. (1) The permutation p: the analysis tries 10,000 random relabellings. With two runs a cell, every possible relabelling (up to 46,656) can be counted, which gives the exact p. exact_p.py does that, written separately from the analysis code. The analysis's p must lie within Monte Carlo error of the exact p. (2) The bootstrap intervals (steps, quality, H3, H4): Excel computes a normal interval with the same variance the rescaled bootstrap targets (the unbiased stratified variance, from the w helper columns on Runs). With two runs a stratum the bootstrap is lumpy, so the two should be close, not equal."),
    ("Summary", None),
]
r = 3
for k, v in lines:
    put(readme, f"A{r}", k, BOLD)
    if v:
        c = put(readme, f"B{r}", v)
        c.alignment = WRAP
        readme.row_dimensions[r].height = max(15, 15 * (len(v) // 120 + 1))
    r += 1
sumrows = [
    ("Runs: read the files right", f'=COUNTIF(Runs!${C["check: analysis read the files right (vs planted)"]}$2:${C["check: analysis read the files right (vs planted)"]}${RN},"ok")', N),
    ("Runs: Excel = analysis", f'=COUNTIF(Runs!${C["check: Excel recompute = analysis"]}$2:${C["check: Excel recompute = analysis"]}${RN},"ok")', N),
    ("Voids: one analysed attempt per slot", f'=COUNTIF(Voids!$J${K1}:$J${KN},"ok")', len(key)),
    ("Voids: counts per model x arm", f'=COUNTIF(Voids!$V$3:$V${VS_LAST},"ok")', len(voids)),
    ("Tests: counts, means, estimates", f'=COUNTIF(Tests!${T["check: counts, means, estimate"]}$2:${T["check: counts, means, estimate"]}${TN},"ok")', len(tests)),
    ("Tests: Newcombe intervals (exact)", f'=COUNTIF(Tests!${T["check: intervals"]}$2:${T["check: intervals"]}${TN},"ok")', sum(1 for t in tests if t["interval"] == "newcombe")),
    ("Tests: bootstrap intervals (close)", f'=COUNTIF(Tests!${T["check: intervals"]}$2:${T["check: intervals"]}${TN},"close")', sum(1 for t in tests if t["interval"] == "bootstrap" and t["analysis"] == "confirmatory")),
    ("Tests: p within Monte Carlo error of exact", f'=COUNTIF(Tests!${T["check: p"]}$2:${T["check: p"]}${TN},"ok")', len(tests)),
    ("Tests: Holm", f'=COUNTIF(Tests!${T["check: Holm"]}$2:${T["check: Holm"]}${TN},"ok")', len(tests)),
    ("Tests: verdicts", f'=COUNTIF(Tests!${T["check: verdict"]}$2:${T["check: verdict"]}${TN},"ok")', len(tests)),
    ("Tests: X", f'=COUNTIF(Tests!${T["check: X"]}$2:${T["check: X"]}${TN},"ok")', len(tests)),
]
put(readme, f"A{r}", "check", BOLD, HEAD_FILL)
put(readme, f"B{r}", "passed / out of", BOLD, HEAD_FILL)
r += 1
for name, f, n in sumrows:
    put(readme, f"A{r}", name)
    put(readme, f"B{r}", f'={f[1:]}&" of {n}"')
    readme.conditional_formatting.add(f"B{r}", CellIsRule(operator="equal", formula=[f'"{n} of {n}"'], fill=OK_FILL))
    r += 1

for sh in wb.worksheets:
    for row_ in sh.iter_rows():
        for cell in row_:
            if cell.font is None or cell.font.name != "Arial":
                f = cell.font
                cell.font = Font(name="Arial", size=f.size or 10, bold=f.bold, color=f.color)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
